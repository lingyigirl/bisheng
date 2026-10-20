"""Pure report rendering. Resources must already be authorized by the caller."""

import base64
import io
from urllib.parse import unquote

import markdown
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageOps

from bisheng.common.utils.markdown_cmpnt.md_to_pdf import sanitize_html_for_pdf

MAX_SOURCE = 2 * 1024 * 1024
MAX_RESOURCES = 64 * 1024 * 1024
CSS = """@page { size: A4; margin: 18mm; }
body { font-family: 'Noto Sans CJK SC','WenQuanYi Zen Hei','Microsoft YaHei',sans-serif;
 font-size: 11pt; line-height: 1.6; overflow-wrap: anywhere; }
h1,h2,h3,h4 { page-break-after: avoid; } img { max-width: 100%; height: auto; }
table { border-collapse: collapse; width: 100%; table-layout: auto; }
td,th { border: 1px solid #888; padding: 5px; overflow-wrap: anywhere; }
thead { display: table-header-group; } tr { break-inside: avoid; }
pre { white-space: pre-wrap; } blockquote { border-left: 3px solid #888; padding-left: 12px; }
"""


def source_tree(source: str):
    if not source.strip() or len(source.encode("utf-8")) > MAX_SOURCE:
        raise ValueError("Report source is empty or exceeds 2 MiB")
    return BeautifulSoup(markdown.markdown(source, extensions=["tables", "fenced_code", "sane_lists"]), "html.parser")


def resource_references(source: str) -> list[str]:
    return list(dict.fromkeys(str(node.get("src", "")) for node in source_tree(source).find_all("img")))


def _prepared(source: str, resources: dict[str, bytes]):
    soup = source_tree(source)
    refs = list(dict.fromkeys(str(node.get("src", "")) for node in soup.find_all("img")))
    if len(refs) > 50 or sum(len(value) for value in resources.values()) > MAX_RESOURCES:
        raise ValueError("Report resource limit exceeded")
    images = {}
    for ref in refs:
        data = resources.get(ref, resources.get(unquote(ref)))
        if data is None:
            raise ValueError(f"Missing report image: {ref}")
        if len(data) > 10 * 1024 * 1024:
            raise ValueError(f"Report image exceeds 10 MiB: {ref}")
        try:
            with Image.open(io.BytesIO(data)) as image:
                if image.width * image.height > 40_000_000:
                    raise ValueError("Image pixel limit exceeded")
                image.verify()
                mime = Image.MIME.get(image.format)
                if mime not in {"image/png", "image/jpeg", "image/gif", "image/webp"}:
                    raise ValueError("Unsupported report image format")
        except (OSError, SyntaxError) as exc:
            raise ValueError(f"Invalid report image: {ref}") from exc
        uri = f"data:{mime};base64,{base64.b64encode(data).decode()}"
        images[uri] = data
        for node in soup.find_all("img", src=ref):
            node["src"] = uri
            node.attrs = {key: value for key, value in node.attrs.items() if key in {"src", "alt", "title"}}
    # Embed authorized resources first; the shared sanitizer still rejects arbitrary URLs.
    clean = BeautifulSoup(sanitize_html_for_pdf(str(soup)), "html.parser")
    if not clean.get_text(strip=True) and not clean.find("img"):
        raise ValueError("Report has no renderable content after sanitization")
    return clean, images


def render_html(source: str, resources: dict[str, bytes]) -> str:
    soup, _ = _prepared(source, resources)
    return f'<!doctype html><html><head><meta charset="utf-8"><style>{CSS}</style></head><body>{soup}</body></html>'


def render_docx(source: str, resources: dict[str, bytes]) -> bytes:
    soup, images = _prepared(source, resources)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.left_margin = section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.font.name = "Arial"
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normalized_images = {}

    def inline(node, paragraph, bold=False, italic=False):
        if isinstance(node, NavigableString):
            run = paragraph.add_run(str(node))
            run.bold, run.italic = bold, italic
        elif node.name == "img":
            data = images.get(node.get("src"))
            if data is None:
                raise ValueError("Unresolved report image")
            with Image.open(io.BytesIO(data)) as image:
                width = min(6.8, image.width / 96)
                if node.get("src") not in normalized_images:
                    stream = io.BytesIO()
                    ImageOps.exif_transpose(image).convert("RGBA").save(stream, format="PNG")
                    normalized_images[node.get("src")] = stream.getvalue()
            paragraph.add_run().add_picture(io.BytesIO(normalized_images[node.get("src")]), width=Inches(width))
        elif node.name == "br":
            paragraph.add_run().add_break()
        else:
            for child in node.children:
                inline(child, paragraph, bold or node.name in {"strong", "b"}, italic or node.name in {"em", "i"})

    def numbering(ordered, start=1):
        root = doc.part.numbering_part.element
        existing = [int(x) for x in root.xpath("./w:abstractNum/@w:abstractNumId")]
        abstract_id = max(existing, default=0) + 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        for tag, value in [
            ("start", str(start)),
            ("numFmt", "decimal" if ordered else "bullet"),
            ("lvlText", "%1." if ordered else "•"),
        ]:
            element = OxmlElement(f"w:{tag}")
            element.set(qn("w:val"), value)
            level.append(element)
        abstract.append(level)
        root.append(abstract)
        num = root.add_num(abstract_id)
        return num.numId

    def blocks(parent, target=doc, depth=0):
        for node in parent.children:
            name = getattr(node, "name", None)
            if not name:
                if str(node).strip():
                    inline(node, target.add_paragraph())
                continue
            if name in {"ol", "ul"}:
                num_id = numbering(name == "ol", int(node.get("start", 1)))
                for li in node.find_all("li", recursive=False):
                    paragraph = None
                    first = True
                    for child in li.children:
                        child_name = getattr(child, "name", None)
                        # Loose lists contain formatting whitespace around <p>.
                        # Rendering it as a run puts the number on an empty line.
                        if isinstance(child, NavigableString) and not child.strip():
                            continue
                        if child_name in {"ol", "ul"}:
                            wrapper = BeautifulSoup("", "html.parser")
                            wrapper.append(child.__copy__())
                            blocks(wrapper, target, depth + 1)
                            paragraph = None
                            continue
                        if paragraph is None or child_name in {"p", "pre", "blockquote"}:
                            paragraph = target.add_paragraph()
                            paragraph.paragraph_format.left_indent = Inches(0.25 * (depth + 1))
                            if first:
                                paragraph.paragraph_format.first_line_indent = Inches(-0.2)
                                props = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                                props.get_or_add_ilvl().val = 0
                                props.get_or_add_numId().val = num_id
                                first = False
                        inline(child, paragraph)
            elif name == "table":
                rows = node.find_all("tr")
                count = max((len(row.find_all(["th", "td"], recursive=False)) for row in rows), default=0)
                if not count:
                    continue
                table = target.add_table(rows=0, cols=count)
                table.style = "Table Grid"
                weights = [8] * count
                for row in rows:
                    for index, cell in enumerate(row.find_all(["th", "td"], recursive=False)):
                        weights[index] = max(weights[index], min(40, len(cell.get_text())))
                usable_width = section.page_width - section.left_margin - section.right_margin
                for column, weight in zip(table.columns, weights, strict=True):
                    column.width = int(usable_width * weight / sum(weights))
                for row in rows:
                    cells = table.add_row().cells
                    row_props = table.rows[-1]._tr.get_or_add_trPr()
                    row_props.append(OxmlElement("w:cantSplit"))
                    for index, cell in enumerate(row.find_all(["th", "td"], recursive=False)):
                        inline(cell, cells[index].paragraphs[0], bold=cell.name == "th")
                        cells[index].paragraphs[0].paragraph_format.space_after = Pt(4)
                    if row.find("th"):
                        header = OxmlElement("w:tblHeader")
                        row_props.append(header)
            elif name in {"p", "pre", "blockquote", "h1", "h2", "h3", "h4", "h5", "h6"}:
                paragraph = target.add_paragraph(style=f"Heading {name[1]}" if name.startswith("h") else None)
                if name == "pre":
                    paragraph.paragraph_format.space_after = Pt(6)
                inline(node, paragraph)
            elif name in {"hr"}:
                target.add_paragraph("―" * 20)
            else:
                blocks(node, target, depth)

    blocks(soup)
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()
