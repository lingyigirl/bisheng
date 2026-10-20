"""Validate actual artifact bytes before advertising successful delivery."""

import csv
import io
import zipfile
from pathlib import PurePosixPath


def validate_artifact(name: str, data: bytes, *, max_expanded_bytes=256 * 1024 * 1024) -> dict:
    if not data or len(data) > 64 * 1024 * 1024:
        raise ValueError("Artifact is empty or exceeds 64 MiB")
    suffix = PurePosixPath(name).suffix.lower()
    manifest = {"bytes": len(data), "type": suffix}
    try:
        if suffix in {".docx", ".xlsx", ".pptx"}:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(x.file_size for x in entries) > max_expanded_bytes:
                    raise ValueError("Artifact expanded ZIP limit exceeded")
                if archive.testzip():
                    raise ValueError("Artifact ZIP checksum mismatch")
        if suffix == ".docx":
            from docx import Document

            doc = Document(io.BytesIO(data))
            if not doc.paragraphs and not doc.tables:
                raise ValueError("Empty Word document")
            manifest.update(images=len(doc.inline_shapes), tables=len(doc.tables))
        elif suffix == ".pdf":
            import fitz

            with fitz.open(stream=data, filetype="pdf") as doc:
                if doc.is_encrypted or not len(doc):
                    raise ValueError("PDF is encrypted or empty")
                if len(doc) > 1000:
                    raise ValueError("PDF exceeds 1000 pages")
                # Force page parsing. Font/layout correctness also needs visual QA.
                image_ids = set()
                image_pages = []
                for page in doc:
                    page.get_text()
                    images = page.get_images()
                    image_ids.update(image[0] for image in images)
                    # get_images() includes shared resources even on pages that
                    # never draw them; image_info inspects actual page content.
                    if len(image_pages) < 3 and page.get_image_info():
                        image_pages.append(page.number + 1)
                manifest["pages"] = len(doc)
                manifest.update(raster_images=len(image_ids), image_pages=image_pages)
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(data), read_only=True)
            try:
                if not workbook.sheetnames:
                    raise ValueError("Spreadsheet has no sheets")
                # Read-only mode defers worksheet XML parsing until iteration.
                # Force it inside the cancellable validation process.
                for sheet in workbook:
                    for _ in sheet.iter_rows(values_only=True):
                        pass
                manifest["sheets"] = len(workbook.sheetnames)
            finally:
                workbook.close()
        elif suffix == ".pptx":
            from pptx import Presentation

            manifest["slides"] = len(Presentation(io.BytesIO(data)).slides)
            if not manifest["slides"]:
                raise ValueError("Presentation has no slides")
        elif suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
            from PIL import Image

            with Image.open(io.BytesIO(data)) as image:
                if image.width * image.height > 40_000_000:
                    raise ValueError("Image pixel limit exceeded")
                manifest.update(width=image.width, height=image.height)
                image.verify()
        elif suffix in {".md", ".markdown", ".html", ".htm", ".csv"}:
            text = data.decode("utf-8-sig")
            if not text.strip():
                raise ValueError("Text artifact is empty")
            if suffix == ".csv":
                manifest["rows"] = sum(1 for _ in csv.reader(io.StringIO(text), strict=True))
        return manifest
    except Exception as exc:
        raise ValueError(f"Invalid {suffix} artifact: {exc}") from exc
