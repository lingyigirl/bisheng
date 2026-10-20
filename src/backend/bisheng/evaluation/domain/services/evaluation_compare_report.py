"""F090: 评测可视化对比报告 — 数据组装（只读）。

本模块自建 report 所需的数据侧逻辑：读结果 CSV（**跳过末尾总行**）、结论引擎、
对比对象条、幻觉率聚合、报告数据组装、两版导出。**不修改** F066 / F053 既有文件；
指标重合处 import 复用其纯函数（见 design 决策 3）。
"""

from __future__ import annotations

import io
import math
from collections import defaultdict, deque
from typing import Any

import pandas as pd
from fastapi import HTTPException
from loguru import logger

from bisheng.evaluation.domain.services.evaluation_metrics import summarize_performance

_NA = "N/A"

# Result-CSV columns surfaced to the report.
_PERCENT_COLUMNS = ("recall", "precision", "F1")
_PLAIN_NUMBER_COLUMNS = ("duration_ms", "input_token", "output_token", "first_token_ms")
_TEXT_COLUMNS = ("question", "ground_truth", "answer", "exec_status")
# F090: hallucination raw counts — present only once the eval chain writes them (T08).
_HALLUCINATION_COLUMNS = ("statements_contradicted_num", "statements_num_answer")

_NA_TOKENS = ("na", "n/a", "nan", "nan%", "none")


def _is_blank(value: Any) -> bool:
    """True for None / NaN / empty-or-whitespace-only text."""
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return not str(value).strip()


def _to_number(value: Any) -> float | None:
    """Plain numeric parse for latency / token cells (independent of the percent parser).

    Deliberately does **not** divide by 100: a pure-int column read by pandas is an
    ``np.int64`` which is not a Python ``int``; a percent parser would turn ``3200``
    into ``32`` (design 坑 2).
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        v = float(value)
        return None if math.isnan(v) else v
    text = str(value).strip()
    if not text or text.lower() in _NA_TOKENS:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _to_percent(value: Any) -> float | None:
    """Parse a percent cell (``"48.49%"``) into a fraction; ``N/A`` / empty -> None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        v = float(value)
        return None if math.isnan(v) else v
    text = str(value).strip()
    if not text or text.lower() in _NA_TOKENS:
        return None
    try:
        return float(text.rstrip("%")) / 100.0
    except ValueError:
        return None


def parse_result_rows(csv_data: Any) -> list[dict[str, Any]]:
    """Read a result CSV into per-question rows, skipping the trailing total row.

    ``build_result_table`` appends a total row whose percent cells are non-empty
    strings (``"58.34%"``), so ``dropna(how="all")`` cannot drop it (design 坑 3).
    We drop on the ``question`` column instead: a real question row always has a
    non-blank question (form rows get a form summary overlaid).

    ``csv_data`` accepts the CSV bytes/text (or a file-like object) — the MinIO
    read is a thin caller-side wrapper so this stays a pure, unit-testable function.
    """
    if hasattr(csv_data, "read"):
        buffer: Any = csv_data.read()
    else:
        buffer = csv_data
    if isinstance(buffer, (bytes, bytearray)):
        buffer = io.BytesIO(buffer)
    elif isinstance(buffer, str):
        buffer = io.StringIO(buffer)
    # utf-8-sig strips the BOM written since 2026-09-06 and still reads legacy files.
    df = pd.read_csv(buffer, encoding="utf-8-sig")

    rows: list[dict[str, Any]] = []
    for _, raw in df.iterrows():
        if _is_blank(raw.get("question")):
            continue  # total row / blank line is not a question
        record: dict[str, Any] = {col: raw.get(col) for col in _TEXT_COLUMNS}
        for col in _PERCENT_COLUMNS:
            record[col] = _to_percent(raw.get(col))
        for col in _PLAIN_NUMBER_COLUMNS:
            record[col] = _to_number(raw.get(col))
        for col in _HALLUCINATION_COLUMNS:
            record[col] = _to_number(raw.get(col))
        status = raw.get("exec_status")
        record["exec_status"] = None if _is_blank(status) else str(status).strip().lower()
        rows.append(record)
    return rows


def compute_success_rate(rows: list[dict[str, Any]]) -> float | None:
    """Success rate over all questions.

    ``None`` (N/A) when no row carries an ``exec_status`` at all — a legacy CSV
    predating F066 must not be reported as 0% (design 坑 3).
    """
    if not rows:
        return None
    statuses = [r.get("exec_status") for r in rows]
    if not any(s in ("success", "failed") for s in statuses):
        return None
    return sum(1 for s in statuses if s == "success") / len(rows)


def compute_hallucination_rate(rows: list[dict[str, Any]]) -> float | None:
    """Content-level hallucination rate: Σcontradicted ÷ Σanswer-statements.

    Deliberately **not** a per-question mean (design 坑 4): a question with many
    statements must weigh more than a one-liner. ``contradicted`` is a subset of
    the answer-only claims, so it is never added to the denominator — the total is
    the answer's own statement count (overlap + answer-only), computed judge-side.

    Rows missing either count (legacy CSV, judge without the 4th bucket, failed
    question → N/A) are skipped from **both** sums; ``None`` when nothing remains
    or the denominator is zero.
    """
    contradicted = 0.0
    total = 0.0
    seen = False
    for row in rows:
        c = row.get("statements_contradicted_num")
        t = row.get("statements_num_answer")
        if not (_is_number(c) and _is_number(t)):
            continue
        seen = True
        contradicted += c
        total += t
    if not seen or total <= 0:
        return None
    return contradicted / total


# ---------------------------------------------------------------------------
# F090: conclusion engine — priority chain + direction normalization + per-type
# tolerance (design 决策 4). Replaces F053's three-metric, same-direction chain.
# ---------------------------------------------------------------------------

# Per-metric tolerance ("按指标类型分别取值", spec AC-04). Percent metrics share one
# value; the count/time metrics are set individually.
_PERCENT_TOLERANCE = 0.01  # 1 percentage point
_TOKEN_TOLERANCE = 1  # 1 token
_LATENCY_TOLERANCE = 50  # 50 ms

# (metric key, Chinese label, higher_is_better, tolerance) in conclusion priority order.
# The key MUST match the metric keys produced by ``aggregate_metrics``
# ("precision" — 准确率 is precision, TP/(TP+FP)), else the top-priority metric is
# silently skipped and the chain starts one level too low.
CONCLUSION_CHAIN: list[tuple[str, str, bool, float]] = [
    ("precision", "准确率", True, _PERCENT_TOLERANCE),
    ("hallucination", "幻觉率", False, _PERCENT_TOLERANCE),
    ("token", "成本(Token)", False, _TOKEN_TOLERANCE),
    ("latency", "延迟", False, _LATENCY_TOLERANCE),
]


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and not math.isnan(float(value))


def decide_conclusion(
    v1: dict[str, Any],
    v2: dict[str, Any],
    unusable_metrics: set[str] | None = None,
) -> dict[str, Any]:
    """Decide which version is better and return the structured analysis.

    ``v1`` / ``v2`` map metric key -> aggregate value (``None`` when unavailable).
    Walks ``CONCLUSION_CHAIN`` in priority order; a metric is only comparable when
    **both** sides have a value, and the first whose normalized delta exceeds its
    tolerance decides. Ties descend; all-tie or all-missing -> ``unable``.

    ``unusable_metrics`` are metrics whose two aggregates were computed from
    **disjoint question sets** (e.g. v1 only has Q1 scored, v2 only Q2). The
    numbers exist but are not comparable, so once such a metric is reached with
    no earlier decider the verdict is ``unable`` — descending to a lower-priority
    metric would silently hide invalid data.

    ``analysis["descent"]`` is True when the decider is not the top-priority metric
    (the report writes the "下探" wording then, spec AC-05). ``delta`` is the raw
    ``v2 - v1``; ``better`` (and the walk) use the direction-normalized comparison.
    """
    unusable = unusable_metrics or set()
    comparisons: list[dict[str, Any]] = []
    conclusion = "unable"
    decider: str | None = None
    unusable_seen: list[str] = []
    blocked = False

    for key, label, higher_is_better, tolerance in CONCLUSION_CHAIN:
        a, b = v1.get(key), v2.get(key)
        if not (_is_number(a) and _is_number(b)):
            continue  # one side missing -> not comparable, skip (spec 边界)
        delta = b - a
        # direction-normalized: positive means v2 is better
        effective = delta if higher_is_better else -delta
        if effective > tolerance:
            better = "v2"
        elif effective < -tolerance:
            better = "v1"
        else:
            better = "tie"
        entry = {"key": key, "label": label, "v1": a, "v2": b, "delta": delta, "better": better}
        comparisons.append(entry)
        if key in unusable:
            entry["unusable"] = True
            unusable_seen.append(key)
            # No decider yet -> the data needed to rank is invalid; never let a
            # lower-priority metric (e.g. token) crown a winner on broken inputs.
            blocked = decider is None
            continue
        if better != "tie" and decider is None and not blocked:
            conclusion = better
            decider = key
            # the first exceeding metric decides; keep collecting the rest for display
    descent = decider is not None and decider != CONCLUSION_CHAIN[0][0]
    return {
        "conclusion": conclusion,
        "analysis": {
            "decider": decider,
            "descent": descent,
            "comparisons": comparisons,
            "unusable_metrics": unusable_seen,
        },
    }


# ---------------------------------------------------------------------------
# F090: 对比对象条 — the two versions' LLM nodes (model + prompts) and which
# changed (design 决策 9). Node diff is delegated to F051's pure
# ``compare_flow_versions`` (reused, not modified).
# ---------------------------------------------------------------------------

_LLM_NODE_TYPE = "llm"
_MODEL_PATH = "model_id"
_PROMPT_PATHS = ("system_prompt", "user_prompt")


def _flatten_node_params(node_data: dict) -> dict[str, Any]:
    """Flatten a node's ``group_params[].params[]`` into ``{key: value}``."""
    flat: dict[str, Any] = {}
    for group in node_data.get("group_params") or []:
        for param in group.get("params") or []:
            key = param.get("key")
            if key is not None:
                flat[key] = param.get("value")
    return flat


def _llm_nodes(data: dict | None) -> list[dict[str, Any]]:
    """The LLM nodes of a ``flow_version.data`` snapshot (prompts live inside them)."""
    result: list[dict[str, Any]] = []
    for node in (data or {}).get("nodes") or []:
        node_data = node.get("data") or {}
        if node_data.get("type") != _LLM_NODE_TYPE:
            continue
        params = _flatten_node_params(node_data)
        result.append(
            {
                "node_id": node.get("id"),
                "node_name": node_data.get("name"),
                "model_id": params.get("model_id"),
                "system_prompt": params.get("system_prompt"),
                "user_prompt": params.get("user_prompt"),
            }
        )
    return result


def _model_name(model_id: Any, model_names: dict[int, str] | None) -> str:
    if model_id is None:
        return _NA
    if model_names and model_id in model_names:
        return model_names[model_id]
    return str(model_id)


def build_object_bar(
    v1_data: dict | None,
    v2_data: dict | None,
    model_names: dict[int, str] | None = None,
) -> dict[str, Any]:
    """Build the 对比对象条: each version's models / prompt nodes + change kind.

    ``change_kind`` is ``"model"`` / ``"prompt"`` / ``"both"`` / ``"none"``.
    A node added/removed counts as a change on both axes (its model *and* prompts
    are effectively new), so a structural edit surfaces as ``"both"`` (多变量).
    """
    # Imported lazily to keep this module importable without the diff util's deps.
    from bisheng.common.utils.flow_version_diff import compare_flow_versions

    v1_nodes = _llm_nodes(v1_data)
    v2_nodes = _llm_nodes(v2_data)

    diff = compare_flow_versions(v1_data or {}, v2_data or {})
    changed_llm = [c for c in diff.get("node_changes", []) if c.get("node_type") == _LLM_NODE_TYPE]

    model_changed = False
    prompt_changed = False
    for change in changed_llm:
        if change.get("change") in ("added", "removed"):
            model_changed = True
            prompt_changed = True
            continue
        paths = [f.get("path", "") for f in change.get("fields", [])]
        if any(_MODEL_PATH in p for p in paths):
            model_changed = True
        if any(any(pp in p for pp in _PROMPT_PATHS) for p in paths):
            prompt_changed = True

    if model_changed and prompt_changed:
        change_kind = "both"
    elif model_changed:
        change_kind = "model"
    elif prompt_changed:
        change_kind = "prompt"
    else:
        change_kind = "none"

    return {
        "models": {
            "v1": [_model_name(n["model_id"], model_names) for n in v1_nodes],
            "v2": [_model_name(n["model_id"], model_names) for n in v2_nodes],
        },
        "prompt_nodes": {
            "v1": [n["node_name"] for n in v1_nodes],
            "v2": [n["node_name"] for n in v2_nodes],
        },
        "change_kind": change_kind,
    }


# ---------------------------------------------------------------------------
# F090: 指标卡 — the six report metrics. ``higher_is_better`` drives the
# frontend's better/worse flip (spec AC-08); "达标" judgment belongs to 6.4 and
# is deliberately absent (spec AC-06).
# ---------------------------------------------------------------------------

METRIC_DEFS: list[tuple[str, str, str, bool]] = [
    ("precision", "准确率", "%", True),
    ("f1", "F1", "%", True),
    ("recall", "召回", "%", True),
    ("hallucination", "幻觉率", "%", False),
    ("token", "成本(Token)", "", False),
    ("latency", "延迟", "ms", False),
]

# The three percent effect metrics share one axis in the grouped bar chart
# (spec AC-07 / design 坑 1); hallucination is % too but stays out (direction).
_GROUPED_BAR_KEYS = ("precision", "f1", "recall")


def _mean(values: list[Any]) -> float | None:
    nums = [v for v in values if _is_number(v)]
    return sum(nums) / len(nums) if nums else None


def aggregate_metrics(rows: list[dict[str, Any]], perf: dict[str, Any] | None = None) -> dict[str, float | None]:
    """Aggregate one version's per-question rows into the six conclusion metrics.

    Effect metrics are the mean of per-question fractions (same rule as the result
    CSV total row); hallucination is content-level Σ/Σ; token is the total token
    sum; latency is the average end-to-end duration. Any missing input -> ``None``.
    """
    perf = perf if perf is not None else summarize_performance(rows)
    return {
        "precision": _mean([r.get("precision") for r in rows]),
        "f1": _mean([r.get("F1") for r in rows]),
        "recall": _mean([r.get("recall") for r in rows]),
        "hallucination": compute_hallucination_rate(rows),
        "token": perf.get("total_token"),
        "latency": perf.get("avg_duration_ms"),
    }


def _delta(v2: Any, v1: Any) -> float | None:
    return v2 - v1 if _is_number(v1) and _is_number(v2) else None


# Metric key -> tolerance used both by the conclusion chain and the per-metric
# direction wording. F1 / recall share the percent tolerance (they are displayed
# only, never walk the chain).
_METRIC_TOLERANCE: dict[str, float] = {
    "precision": _PERCENT_TOLERANCE,
    "f1": _PERCENT_TOLERANCE,
    "recall": _PERCENT_TOLERANCE,
    "hallucination": _PERCENT_TOLERANCE,
    "token": _TOKEN_TOLERANCE,
    "latency": _LATENCY_TOLERANCE,
}
_CHAIN_KEYS = {key for key, _label, _hib, _tol in CONCLUSION_CHAIN}
_PERCENT_METRICS = {"precision", "f1", "recall", "hallucination"}


def _fmt_metric_value(key: str, value: Any) -> str:
    if not _is_number(value):
        return _NA
    if key in _PERCENT_METRICS:
        return f"{value * 100:.2f}%"
    if key == "token":
        return f"{round(value)}"
    return f"{value:.0f} ms"


def _fmt_metric_delta(key: str, delta: Any) -> str:
    if not _is_number(delta):
        return _NA
    if key in _PERCENT_METRICS:
        return f"{delta * 100:+.2f} 个百分点"
    if key == "token":
        return f"{delta:+.0f} token"
    return f"{delta:+.0f} ms"


def _fmt_tolerance(key: str) -> str:
    if key in _PERCENT_METRICS:
        return "1 个百分点"
    if key == "token":
        return "1 token"
    return "50 ms"


def _completeness_lines(per_question: list[dict[str, Any]], labels: tuple[str, str]) -> list[str]:
    """Point out every question whose six metrics are incomplete and why.

    Three distinct causes are spelled out: execution failure (with the row's
    reason), a successful run whose judge output was incomplete (effect metrics /
    hallucination counts missing), and a question present on only one side.
    """
    issues: list[str] = []
    for idx, entry in enumerate(per_question, start=1):
        for side_key, side_label in (("v1", labels[0]), ("v2", labels[1])):
            cell = entry.get(side_key)
            if cell is None:
                issues.append(f"第 {idx} 题 {side_label} 没有结果，无法参与对比")
                continue
            missing = [label for key, label, _unit, _hib in METRIC_DEFS if not _is_number(cell.get(key))]
            if cell.get("exec_status") == "failed":
                reason = cell.get("fail_reason") or "未记录失败原因"
                suffix = f"，缺少{'、'.join(missing)}" if missing else ""
                issues.append(f"第 {idx} 题 {side_label} 执行失败（{reason}）{suffix}")
            elif missing:
                issues.append(f"第 {idx} 题 {side_label} 执行成功但判分结果不完整，缺少{'、'.join(missing)}")

    if not issues:
        if not per_question:
            return ["数据完整性：没有可对比的题目。"]
        return [f"数据完整性：共 {len(per_question)} 题，两版全部执行成功，六项指标均有数据。"]
    lines = ["数据完整性："] + [f"· {issue}" for issue in issues]
    return lines


def _disjoint_metric_keys(per_question: list[dict[str, Any]]) -> set[str]:
    """Metrics whose two aggregates have **no question in common**.

    Aggregation is per-metric, so a side can end up with a value from Q1 while
    the other side's value comes from Q2. Such a pair is not a comparison of the
    same work — the conclusion engine must refuse to rank on it.
    """
    disjoint: set[str] = set()
    for key, _label, _unit, _hib in METRIC_DEFS:
        questions1 = {
            idx for idx, entry in enumerate(per_question, start=1) if _is_number((entry.get("v1") or {}).get(key))
        }
        questions2 = {
            idx for idx, entry in enumerate(per_question, start=1) if _is_number((entry.get("v2") or {}).get(key))
        }
        if questions1 and questions2 and questions1.isdisjoint(questions2):
            disjoint.add(key)
    return disjoint


def _cross_question_lines(per_question: list[dict[str, Any]], labels: tuple[str, str]) -> list[str]:
    """Warn when a metric's aggregate is averaged over different question sets.

    Aggregation skips rows with a missing value per metric, so the two versions'
    means/sums can be computed from different questions. That is legitimate but
    makes a side-by-side number less meaningful — say which questions fed each.
    """
    caveats: list[str] = []
    for key, label, _unit, _hib in METRIC_DEFS:
        questions1 = [
            idx for idx, entry in enumerate(per_question, start=1) if _is_number((entry.get("v1") or {}).get(key))
        ]
        questions2 = [
            idx for idx, entry in enumerate(per_question, start=1) if _is_number((entry.get("v2") or {}).get(key))
        ]
        if questions1 and questions2 and questions1 != questions2:
            first = "、".join(str(i) for i in questions1)
            second = "、".join(str(i) for i in questions2)
            caveats.append(f"{label}：{labels[0]}来自第 {first} 题，{labels[1]}来自第 {second} 题")
    if not caveats:
        return []
    return ["跨题提示：部分指标由不同题目聚合得出，数值直接比较需谨慎。"] + [f"· {c}" for c in caveats]


def build_analysis(
    decision: dict[str, Any],
    metrics: list[dict[str, Any]] | None = None,
    per_question: list[dict[str, Any]] | None = None,
    labels: tuple[str, str] = ("旧版本", "新版本"),
) -> str:
    """Template-generated multi-dimensional 对比分析 (spec AC-05) — no LLM.

    Four sections, joined by newlines (the report page renders them with
    ``whitespace-pre-line``):
      1. 结论依据 — the priority chain, the deciding metric and its margin.
      2. 指标对比 — one-line direction summary plus which metrics carry the
         verdict; exact values live in the six metric cards, so they are not
         repeated here.
      3. 数据完整性 — per-question N/A causes (execution failure / judge output
         incomplete / missing on one side).
      4. 跨题提示 — when an aggregate was computed from different questions.

    ``metrics`` / ``per_question`` are the payloads already built for the report;
    when omitted the metric lines fall back to the decision's comparisons.
    """
    info = decision.get("analysis") or {}
    decider = info.get("decider")
    conclusion = decision.get("conclusion")
    comparisons = {c["key"]: c for c in info.get("comparisons") or []}
    unusable_seen = info.get("unusable_metrics") or []
    unusable_keys = _disjoint_metric_keys(per_question) if per_question is not None else set()

    values: dict[str, dict[str, Any]] = {m["key"]: m for m in (metrics or [])}
    for key, comparison in comparisons.items():
        values.setdefault(key, comparison)

    chain_text = " → ".join(label for _key, label, _hib, _tol in CONCLUSION_CHAIN)
    if not decider or conclusion == "unable":
        if unusable_seen:
            unusable_labels = "、".join(label for key, label, _hib, _tol in CONCLUSION_CHAIN if key in unusable_seen)
            head = (
                f"结论依据：按 {chain_text} 的优先级依次比较，"
                f"{unusable_labels}的两版数值分别来自不同题目，不构成同一批题目的对比，无法判定优劣。"
            )
        else:
            head = f"结论依据：按 {chain_text} 的优先级依次比较，两版没有超过各自容差的可比差异，无法判定优劣。"
    else:
        comparison = comparisons.get(decider) or {}
        label = comparison.get("label") or next((l for k, l, _h, _t in CONCLUSION_CHAIN if k == decider), decider)
        winner = labels[1] if conclusion == "v2" else labels[0]
        delta = _delta(comparison.get("v2"), comparison.get("v1"))
        margin = (
            f"{_fmt_metric_value(decider, comparison.get('v1'))} → "
            f"{_fmt_metric_value(decider, comparison.get('v2'))}，"
            f"变化 {_fmt_metric_delta(decider, delta)}，超出 {_fmt_tolerance(decider)}容差"
        )
        # "下探" wording: only claim a tie for the metrics that really tied.
        tied: list[str] = []
        missing: list[str] = []
        for key, chain_label, _hib, _tol in CONCLUSION_CHAIN:
            if key == decider:
                break
            if key not in comparisons:
                missing.append(chain_label)
            elif comparisons[key].get("better") == "tie":
                tied.append(chain_label)
        if tied:
            prefix = f"{'、'.join(tied)}打平后，"
        elif missing:
            prefix = f"{'、'.join(missing)}缺少数据，"
        else:
            prefix = ""
        head = f"结论依据：按 {chain_text} 的优先级依次比较，{prefix}{winner}的{label}更优（{margin}）。"

    # When a higher-priority metric is unusable the verdict stays "unable", so the
    # remaining chain metrics must not read as if they carried the decision.
    chain_blocked = conclusion == "unable" and bool(unusable_seen)

    # The report page renders all six values in metric cards right below this
    # text (and the exports carry a metric table), so repeating every value and
    # delta line-by-line here was redundant. Summarize direction only.
    grouped: dict[str, list[str]] = {"v1": [], "v2": [], "tie": [], "na": [], "unusable": []}
    for key, label, _unit, higher_is_better in METRIC_DEFS:
        entry = values.get(key)
        v1, v2 = (entry or {}).get("v1"), (entry or {}).get("v2")
        if not (_is_number(v1) and _is_number(v2)):
            grouped["na"].append(label)
            continue
        if key in unusable_keys:
            grouped["unusable"].append(label)
            continue
        effective = (v2 - v1) if higher_is_better else (v1 - v2)
        tolerance = _METRIC_TOLERANCE[key]
        if effective > tolerance:
            grouped["v2"].append(label)
        elif effective < -tolerance:
            grouped["v1"].append(label)
        else:
            grouped["tie"].append(label)

    def _bucket(count_key: str, wording: str) -> str:
        names = grouped[count_key]
        return f"{wording} {len(names)} 项（{'、'.join(names)}）" if names else ""

    summary_parts = [
        part
        for part in (
            _bucket("v1", f"{labels[0]}更优"),
            _bucket("v2", f"{labels[1]}更优"),
            _bucket("tie", "基本持平"),
            _bucket("na", "缺少单侧数据无法比较"),
            (
                f"数据不可比 {len(grouped['unusable'])} 项（{'、'.join(grouped['unusable'])}：两版数值来自不同题目）"
                if grouped["unusable"]
                else ""
            ),
        )
        if part
    ]
    summary_line = f"指标对比：{len(METRIC_DEFS)} 项指标中，{'，'.join(summary_parts)}。"

    chain_labels = " → ".join(label for _key, label, _hib, _tol in CONCLUSION_CHAIN)
    reference_labels = "、".join(label for key, label, _hib, _tol in METRIC_DEFS if key not in _CHAIN_KEYS)
    role_line = f"参与结论判定：{chain_labels}；{reference_labels}仅供参考，不参与结论"
    if chain_blocked:
        role_line += "。高优先级指标数据不可比，结论未以其为依据"
    role_line += "。"

    lines = [head, summary_line, role_line]
    if per_question is not None:
        lines.extend(_completeness_lines(per_question, labels))
        lines.extend(_cross_question_lines(per_question, labels))
    return "\n".join(lines)


def _row_token(row: dict[str, Any]) -> int | None:
    values = [v for v in (row.get("input_token"), row.get("output_token")) if _is_number(v)]
    return int(sum(values)) if values else None


def _row_hallucination(row: dict[str, Any]) -> float | None:
    contradicted, total = row.get("statements_contradicted_num"), row.get("statements_num_answer")
    if _is_number(contradicted) and _is_number(total) and total > 0:
        return contradicted / total
    return None


def _side_entry(row: dict[str, Any] | None) -> dict[str, Any] | None:
    """One question's six-metric cell for one version (``None`` if unpaired)."""
    if row is None:
        return None
    entry: dict[str, Any] = {
        "precision": row.get("precision"),
        "f1": row.get("F1"),
        "recall": row.get("recall"),
        "hallucination": _row_hallucination(row),
        "token": _row_token(row),
        "latency": row.get("duration_ms"),
        "exec_status": row.get("exec_status"),
    }
    if entry["exec_status"] == "failed":
        entry["fail_reason"] = row.get("answer")
    return entry


def _per_question_entry(question: Any, r1: dict | None, r2: dict | None) -> dict[str, Any]:
    e1, e2 = _side_entry(r1), _side_entry(r2)
    entry: dict[str, Any] = {"question": question, "v1": e1, "v2": e2}
    reasons = [
        str(e["fail_reason"]) for e in (e1, e2) if e and e.get("exec_status") == "failed" and e.get("fail_reason")
    ]
    if reasons:
        entry["failed"] = True
        entry["fail_reason"] = " / ".join(dict.fromkeys(reasons))
    return entry


def build_per_question(v1_rows: list[dict[str, Any]], v2_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Pair the two versions' questions and carry **all six** metrics per side.

    Pairing is by question text via F053's ``_align_questions`` (same dataset ->
    positional fast path); a question on only one side pairs with ``None`` on the
    other (spec AC-08, design 坑 12).
    """
    from bisheng.evaluation.domain.services.evaluation_service import _align_questions

    a = [r for r in v1_rows if not _is_blank(r.get("question"))]
    b = [r for r in v2_rows if not _is_blank(r.get("question"))]
    pairs = _align_questions([r.get("question") for r in a], [r.get("question") for r in b])

    index1: dict[str, deque] = defaultdict(deque)
    for r in a:
        index1[str(r.get("question")).strip()].append(r)
    index2: dict[str, deque] = defaultdict(deque)
    for r in b:
        index2[str(r.get("question")).strip()].append(r)

    result: list[dict[str, Any]] = []
    for x, y in pairs:
        r1 = index1[str(x).strip()].popleft() if x is not None and index1[str(x).strip()] else None
        r2 = index2[str(y).strip()].popleft() if y is not None and index2[str(y).strip()] else None
        result.append(_per_question_entry(x if x is not None else y, r1, r2))
    return result


def _evaluation_meta(side: dict[str, Any], perf: dict[str, Any]) -> dict[str, Any]:
    """The两版元信息 block: identity + the title statistics (spec AC-09/AC-10)."""
    question_count = perf.get("question_count") or 0
    success_count = perf.get("success_count") or 0
    return {
        "id": side.get("id"),
        "version": side.get("version"),
        "version_name": side.get("version_name"),
        "status": side.get("status"),
        "result_score": side.get("result_score"),
        "description": side.get("description"),
        "question_count": question_count,
        "success_count": success_count,
        "failed_count": question_count - success_count,
        "success_rate": perf.get("success_rate"),
    }


def build_report_data(
    side1: dict[str, Any],
    side2: dict[str, Any],
    model_names: dict[int, str] | None = None,
) -> dict[str, Any]:
    """Assemble the report payload from two already-loaded sides (v1, v2).

    Pure (no DB/MinIO): the caller loads each side via ``load_compare_side``. This
    is the single data source both the report page and the export consume
    (design 决策 2).
    """
    rows1, rows2 = side1.get("rows") or [], side2.get("rows") or []
    perf1, perf2 = summarize_performance(rows1), summarize_performance(rows2)
    metrics1, metrics2 = aggregate_metrics(rows1, perf1), aggregate_metrics(rows2, perf2)

    # Per-question pairing is needed first: a metric whose two aggregates have no
    # question in common must not decide the conclusion (see decide_conclusion).
    per_question = build_per_question(rows1, rows2)
    decision = decide_conclusion(metrics1, metrics2, _disjoint_metric_keys(per_question))

    object_bar = build_object_bar(side1.get("flow_data"), side2.get("flow_data"), model_names)
    object_bar["app_name"] = side1.get("app_name") or side2.get("app_name")
    object_bar["data_set"] = side1.get("data_set") or side2.get("data_set")

    metrics = [
        {
            "key": key,
            "label": label,
            "unit": unit,
            "higher_is_better": higher_is_better,
            "v1": metrics1.get(key),
            "v2": metrics2.get(key),
            "delta": _delta(metrics2.get(key), metrics1.get(key)),
        }
        for key, label, unit, higher_is_better in METRIC_DEFS
    ]

    grouped_bar = [m for m in metrics if m["key"] in _GROUPED_BAR_KEYS]
    # The per-question detail built above is fed to the analysis text, so the
    # 数据完整性/跨题提示 sections describe exactly what the table below shows.
    labels = (_side_label(side1), _side_label(side2))

    return {
        "evaluations": [_evaluation_meta(side1, perf1), _evaluation_meta(side2, perf2)],
        "object_bar": object_bar,
        "conclusion": decision["conclusion"],
        "analysis": build_analysis(decision, metrics, per_question, labels),
        "metrics": metrics,
        "charts": {"grouped_bar": grouped_bar},
        "per_question": per_question,
    }


def _side_label(side: dict[str, Any]) -> str:
    """Human-readable version label for the analysis text (name first, then version)."""
    name = side.get("version_name")
    if name:
        return str(name)
    version = side.get("version")
    return f"版本 {version}" if version is not None else "本版本"


def load_compare_side(evaluation: Any) -> dict[str, Any]:
    """Load one version's rows (result CSV) + flow snapshot + display names.

    Read-only: reuses F066's app-meta resolver and the F051 flow snapshot; the CSV
    reader is this feature's own (skips the trailing total row)."""
    from bisheng.database.models.flow_version import FlowVersionDao
    from bisheng.evaluation.domain.services.evaluation_report import _resolve_app_meta
    from bisheng.evaluation.domain.services.evaluation_service import EvaluationService

    if not evaluation.result_file_path:
        raise HTTPException(status_code=400, detail="result csv unreadable")
    raw = EvaluationService.read_csv_file(evaluation.result_file_path)
    if raw is None:
        raise HTTPException(status_code=400, detail="result csv unreadable")
    rows = parse_result_rows(raw)

    version_name = None
    flow_data = None
    if evaluation.version:
        versions = FlowVersionDao.get_list_by_ids(ids=[evaluation.version])
        if versions:
            version_name = versions[0].name
            flow_data = versions[0].data
    app_name, _app_type, meta_version = _resolve_app_meta(evaluation)
    if version_name is None and meta_version not in (None, "--"):
        version_name = meta_version

    return {
        "id": evaluation.id,
        "version": evaluation.version,
        "version_name": version_name,
        "app_name": app_name,
        "data_set": evaluation.file_name,
        "status": evaluation.status,
        "result_score": evaluation.result_score,
        "description": evaluation.description,
        "flow_data": flow_data,
        "rows": rows,
    }


def _collect_model_names(flow_datas: list[dict | None]) -> dict[int, str]:
    """Resolve the LLM model ids seen in both flow snapshots to display names."""
    from bisheng.llm.domain.models.llm_server import LLMModel

    ids = {
        node["model_id"] for data in flow_datas for node in _llm_nodes(data) if isinstance(node.get("model_id"), int)
    }
    if not ids:
        return {}
    try:
        models = LLMModel.get_model_by_ids(list(ids))
    except Exception as e:  # a missing name must not break the report
        logger.warning("compare report model name resolve failed: {}", e)
        return {}
    return {m.id: m.name for m in models}


def build_compare_report(group_id: str, user_id: int) -> dict[str, Any]:
    """Load an owner-scoped compare group and assemble its report data.

    Raises HTTPException 404 when the group is not the user's (or absent) and 400
    when its two members are not the same agent (spec AC-15).
    """
    from bisheng.evaluation.domain.repositories.evaluation_repository import EvaluationRepository

    evaluations = EvaluationRepository.get_compare_group(group_id, user_id)
    if len(evaluations) != 2:
        raise HTTPException(status_code=404, detail="compare group not found")
    evaluations = sorted(evaluations, key=lambda e: (e.version is None, e.version or 0))
    v1, v2 = evaluations
    if str(v1.unique_id) != str(v2.unique_id):
        raise HTTPException(status_code=400, detail="compare group agents mismatch")

    side1, side2 = load_compare_side(v1), load_compare_side(v2)
    model_names = _collect_model_names([side1.get("flow_data"), side2.get("flow_data")])
    data = build_report_data(side1, side2, model_names=model_names)
    data["group_id"] = group_id
    return data


# ---------------------------------------------------------------------------
# F090: 两版对比导出（xlsx）— reuses F066's openpyxl sheet/cell approach and N/A
# convention; the two-version layout is this feature's own (design 决策 7). The
# workbook is generated from the same report data as the page, so it is always
# full (all six metrics per question) — never a function of a selected metric.
# ---------------------------------------------------------------------------

_REPORT_DIR = "evaluation/report"
_XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_STATUS_LABELS = {"success": "成功", "failed": "失败"}
# Short labels for the (wide) detail sheet; the aggregated card uses METRIC_DEFS.
_EXPORT_METRIC_LABELS: list[tuple[str, str]] = [
    ("precision", "准确率"),
    ("f1", "F1"),
    ("recall", "召回"),
    ("hallucination", "幻觉率"),
    ("token", "成本"),
    ("latency", "延迟"),
]
_METRIC_UNITS = {key: unit for key, _label, unit, _hib in METRIC_DEFS}


def _version_label(meta: dict[str, Any]) -> str:
    return meta.get("version_name") or (f"v{meta.get('version')}" if meta.get("version") is not None else "v?")


def _export_cell(key: str, value: Any) -> str:
    """Render one metric value for the sheet (``N/A`` when absent, F066 convention)."""
    if value is None or isinstance(value, bool):
        return _NA
    if _METRIC_UNITS.get(key) == "%":
        return f"{value * 100:.2f}%"
    return f"{value:.2f}".rstrip("0").rstrip(".") if isinstance(value, float) else str(value)


def _detail_frame(per_question: list[dict[str, Any]], label1: str, label2: str) -> pd.DataFrame:
    data: list[dict[str, Any]] = []
    for entry in per_question:
        record: dict[str, Any] = {"题目": entry.get("question")}
        for label, side_key in ((label1, "v1"), (label2, "v2")):
            side = entry.get(side_key) or {}
            record[f"状态({label})"] = _STATUS_LABELS.get(side.get("exec_status") or "", _NA)
            for key, metric_label in _EXPORT_METRIC_LABELS:
                record[f"{metric_label}({label})"] = _export_cell(key, side.get(key))
        data.append(record)
    columns = ["题目"] + [
        f"{header}({label})"
        for label in (label1, label2)
        for header in ["状态"] + [metric_label for _key, metric_label in _EXPORT_METRIC_LABELS]
    ]
    return pd.DataFrame(data, columns=columns)


def build_export_xlsx(data: dict[str, Any]) -> bytes:
    """Render the two-version comparison workbook (汇总 + 逐题明细) into bytes."""
    evaluations = data.get("evaluations") or []
    label1 = _version_label(evaluations[0]) if len(evaluations) > 0 else "v1"
    label2 = _version_label(evaluations[1]) if len(evaluations) > 1 else "v2"

    summary_rows = [
        [
            metric["label"],
            _export_cell(metric["key"], metric["v1"]),
            _export_cell(metric["key"], metric["v2"]),
            _export_cell(metric["key"], metric["delta"]),
        ]
        for metric in data.get("metrics") or []
    ]
    summary = pd.DataFrame(summary_rows, columns=["指标", label1, label2, "差值"])
    detail = _detail_frame(data.get("per_question") or [], label1, label2)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="汇总", index=False)
        detail.to_excel(writer, sheet_name="逐题明细", index=False)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# F090 (2026-10-19): full report export (HTML / PDF). The user asked the exported
# file to be a forwardable report, not just the data workbook. HTML is
# self-contained (inline CSS + SVG charts); PDF is rendered by the project's
# existing Playwright/Chromium stack (same approach as conversation export).
# ---------------------------------------------------------------------------

EXPORT_CONTENT_TYPES = {
    "xlsx": _XLSX_CONTENT_TYPE,
    "html": "text/html; charset=utf-8",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

_CHANGE_KIND_LABELS = {
    "model": "模型",
    "prompt": "提示词",
    "both": "模型与提示词",
    "none": "无变化",
}

_EXPORT_HTML_CSS = """
@page { size: A4; margin: 12mm 10mm; }
* { box-sizing: border-box; }
body { margin: 0; color: #1f2937; font-size: 12px; line-height: 1.6;
       font-family: "PingFang SC", "Microsoft YaHei", "WenQuanYi Zen Hei",
                    "Noto Sans CJK SC", "Source Han Sans SC", sans-serif; }
h1 { margin: 0 0 2px; font-size: 22px; }
h2 { margin: 18px 0 8px; padding-left: 8px; font-size: 15px;
     border-left: 4px solid #2563eb; }
.meta, .small, .muted { color: #6b7280; font-size: 11px; }
.object-bar { border: 1px solid #bfdbfe; background: #eff6ff; border-radius: 6px;
              padding: 10px 12px; }
.object-line { margin-bottom: 8px; }
.version-grid { display: grid; grid-template-columns: 1fr 1fr; gap: 8px; }
.version-card { border: 1px solid #bfdbfe; border-radius: 4px; background: #ffffff;
                padding: 6px 8px; }
.version-name { font-weight: 600; }
.tag { display: inline-block; margin-left: 6px; padding: 0 6px; border-radius: 4px;
       background: #fef3c7; color: #92400e; font-size: 10px; }
.warning { margin-top: 8px; padding: 6px 8px; border: 1px solid #fcd34d;
           border-radius: 4px; background: #fffbeb; color: #92400e; font-size: 11px; }
.conclusion-box { border: 1px solid #e5e7eb; border-radius: 6px; padding: 8px 10px; }
.verdict { font-size: 14px; font-weight: 600; }
.analysis { margin-top: 4px; white-space: pre-line; color: #4b5563; }
.cards { display: grid; grid-template-columns: repeat(3, 1fr); gap: 8px; }
.card { border: 1px solid #e5e7eb; border-radius: 6px; padding: 8px 10px; }
.card-title { margin-bottom: 4px; font-size: 11px; font-weight: 600; color: #6b7280; }
.row { display: flex; justify-content: space-between; gap: 8px; }
.delta { margin-top: 2px; text-align: right; font-size: 10px; color: #6b7280; }
.chart { border: 1px solid #e5e7eb; border-radius: 6px; padding: 4px; background: #ffffff; }
.chart-empty { padding: 12px; color: #9ca3af; font-size: 11px; }
.line-grid { display: grid; grid-template-columns: 1fr 1fr; gap: 8px; }
.detail-title { margin: 10px 0 4px; font-weight: 600; }
table { width: 100%; border-collapse: collapse; font-size: 10px; }
th, td { border: 1px solid #e5e7eb; padding: 4px 5px; text-align: left; vertical-align: top; }
th { background: #f3f4f6; font-weight: 600; }
tr { page-break-inside: avoid; }
.question { white-space: pre-wrap; word-break: break-word; }
.fail { color: #dc2626; font-size: 9px; }
"""


def _esc(value: Any) -> str:
    """HTML-escape a dynamic value (app / dataset / question text / analysis)."""
    from html import escape

    return escape(_NA if value is None else str(value), quote=True)


def _axis_value(key: str, value: float) -> str:
    if key in _PERCENT_METRICS:
        return f"{value * 100:.0f}%"
    if value >= 1000:
        return f"{value / 1000:.1f}k"
    return f"{value:.0f}"


def _series_values(per_question: list[dict[str, Any]], side_key: str, key: str) -> list[Any]:
    values: list[Any] = []
    for entry in per_question:
        side = entry.get(side_key) or {}
        value = side.get(key)
        values.append(value if _is_number(value) else None)
    return values


def _line_paths(values: list[Any], color: str, x_at, y_at) -> str:
    segments: list[list[tuple[int, float]]] = []
    current: list[tuple[int, float]] = []
    for index, value in enumerate(values):
        if _is_number(value):
            current.append((index, value))
        elif current:
            segments.append(current)
            current = []
    if current:
        segments.append(current)

    parts: list[str] = []
    for segment in segments:
        if len(segment) >= 2:
            points = " ".join(f"{x_at(i):.1f},{y_at(v):.1f}" for i, v in segment)
            parts.append(f'<polyline points="{points}" fill="none" stroke="{color}" stroke-width="2"/>')
        for i, v in segment:
            parts.append(f'<circle cx="{x_at(i):.1f}" cy="{y_at(v):.1f}" r="2.5" fill="{color}"/>')
    return "".join(parts)


def _svg_line_chart(
    per_question: list[dict[str, Any]],
    key: str,
    label: str,
    version_label1: str,
    version_label2: str,
) -> str:
    """One metric's per-question trend (both versions on the same axis)."""
    width, height = 360, 186
    left, right, top, bottom = 44, 10, 34, 22
    plot_w = width - left - right
    plot_h = height - top - bottom
    n = len(per_question)
    if n == 0:
        return '<div class="chart chart-empty">无逐题数据</div>'

    values1 = _series_values(per_question, "v1", key)
    values2 = _series_values(per_question, "v2", key)
    numbers = [v for v in values1 + values2 if _is_number(v)]
    vmax = max(numbers) if numbers else 0
    vmax = vmax * 1.15 if vmax > 0 else 1.0

    def x_at(index: int) -> float:
        return left + (plot_w * index / (n - 1) if n > 1 else plot_w / 2)

    def y_at(value: float) -> float:
        return top + plot_h * (1 - value / vmax)

    grid: list[str] = []
    for tick in (0.0, 0.5, 1.0):
        y = y_at(vmax * tick)
        grid.append(
            f'<line x1="{left}" y1="{y:.1f}" x2="{width - right}" y2="{y:.1f}" stroke="#e5e7eb" stroke-width="1"/>'
        )
        grid.append(
            f'<text x="{left - 6}" y="{y + 3:.1f}" text-anchor="end" font-size="9" fill="#6b7280">'
            f"{_axis_value(key, vmax * tick)}</text>"
        )
    step = max(1, math.ceil(n / 10))
    for index in range(n):
        if index % step == 0 or index == n - 1:
            grid.append(
                f'<text x="{x_at(index):.1f}" y="{height - bottom + 13}" text-anchor="middle" '
                f'font-size="9" fill="#6b7280">{index + 1}</text>'
            )

    legend = (
        f'<rect x="{left}" y="18" width="9" height="9" fill="#94a3b8"/>'
        f'<text x="{left + 13}" y="26" font-size="9" fill="#374151">{_esc(version_label1)}</text>'
        f'<rect x="{left + 150}" y="18" width="9" height="9" fill="#2563eb"/>'
        f'<text x="{left + 163}" y="26" font-size="9" fill="#374151">{_esc(version_label2)}</text>'
    )
    paths = _line_paths(values1, "#94a3b8", x_at, y_at) + _line_paths(values2, "#2563eb", x_at, y_at)
    return (
        f'<div class="chart"><svg viewBox="0 0 {width} {height}" width="100%">'
        f'<text x="{left}" y="12" font-size="11" font-weight="600" fill="#111827">{_esc(label)}</text>'
        f"{legend}{''.join(grid)}{paths}</svg></div>"
    )


def _svg_grouped_bar(
    grouped: list[dict[str, Any]],
    version_label1: str,
    version_label2: str,
) -> str:
    """Grouped bars for the three percentage effect metrics (one shared axis)."""
    width, height = 700, 230
    left, right, top, bottom = 48, 12, 34, 34
    plot_w = width - left - right
    plot_h = height - top - bottom
    if not grouped:
        return '<div class="chart chart-empty">无指标数据</div>'

    group_w = plot_w / len(grouped)
    bar_w = min(54.0, group_w * 0.26)
    parts: list[str] = []
    for tick in (0.0, 0.25, 0.5, 0.75, 1.0):
        y = top + plot_h * (1 - tick)
        parts.append(
            f'<line x1="{left}" y1="{y:.1f}" x2="{width - right}" y2="{y:.1f}" stroke="#e5e7eb" stroke-width="1"/>'
        )
        parts.append(
            f'<text x="{left - 6}" y="{y + 3:.1f}" text-anchor="end" font-size="9" fill="#6b7280">'
            f"{_axis_value('precision', tick)}</text>"
        )
    for index, metric in enumerate(grouped):
        center = left + group_w * (index + 0.5)
        parts.append(
            f'<text x="{center:.1f}" y="{height - bottom + 16}" text-anchor="middle" '
            f'font-size="10" fill="#374151">{_esc(metric.get("label"))}</text>'
        )
        for offset, value, color in (
            (-bar_w / 2 - 3, metric.get("v1"), "#94a3b8"),
            (bar_w / 2 + 3, metric.get("v2"), "#2563eb"),
        ):
            if not _is_number(value):
                continue
            bar_h = plot_h * max(0.0, min(1.0, float(value)))
            x = center + offset - bar_w / 2
            y = top + plot_h - bar_h
            parts.append(
                f'<rect x="{x:.1f}" y="{y:.1f}" width="{bar_w:.1f}" height="{bar_h:.1f}" fill="{color}" rx="1"/>'
            )
            parts.append(
                f'<text x="{center + offset:.1f}" y="{y - 3:.1f}" text-anchor="middle" '
                f'font-size="9" fill="#374151">{float(value) * 100:.1f}%</text>'
            )
    legend = (
        f'<rect x="{left}" y="16" width="10" height="10" fill="#94a3b8"/>'
        f'<text x="{left + 14}" y="25" font-size="10" fill="#374151">{_esc(version_label1)}</text>'
        f'<rect x="{left + 150}" y="16" width="10" height="10" fill="#2563eb"/>'
        f'<text x="{left + 164}" y="25" font-size="10" fill="#374151">{_esc(version_label2)}</text>'
    )
    return f'<div class="chart"><svg viewBox="0 0 {width} {height}" width="100%">{legend}{"".join(parts)}</svg></div>'


def _html_detail_table(per_question: list[dict[str, Any]], side_key: str, label: str) -> str:
    """One version's per-question metrics table (all six metrics, F066 N/A口径)."""
    headers = "".join(f"<th>{_esc(label)}</th>" for _key, label in _EXPORT_METRIC_LABELS)
    rows: list[str] = []
    for index, entry in enumerate(per_question, start=1):
        side = entry.get(side_key) or {}
        status = _STATUS_LABELS.get(side.get("exec_status") or "", _NA)
        note = ""
        if entry.get("failed"):
            reason = entry.get("fail_reason")
            note = f'<div class="fail">失败：{_esc(reason)}</div>' if reason else '<div class="fail">失败</div>'
        cells = "".join(f"<td>{_export_cell(key, side.get(key))}</td>" for key, _label in _EXPORT_METRIC_LABELS)
        rows.append(
            f'<tr><td>{index}</td><td class="question">{_esc(entry.get("question"))}{note}</td>'
            f"<td>{status}</td>{cells}</tr>"
        )
    return (
        f'<div class="detail-title">{_esc(label)}</div>'
        "<table><thead><tr><th>#</th><th>题目</th><th>状态</th>"
        f"{headers}</tr></thead><tbody>{''.join(rows)}</tbody></table>"
    )


def build_export_html(data: dict[str, Any]) -> str:
    """Render the full comparison report as a single self-contained HTML document."""
    from datetime import datetime

    evaluations = data.get("evaluations") or []
    label1 = _version_label(evaluations[0]) if len(evaluations) > 0 else "v1"
    label2 = _version_label(evaluations[1]) if len(evaluations) > 1 else "v2"

    object_bar = data.get("object_bar") or {}
    change_kind = object_bar.get("change_kind") or "none"
    change_label = _CHANGE_KIND_LABELS.get(change_kind, change_kind)
    tag = '<span class="tag">多变量 · 不可单独归因</span>' if change_kind == "both" else ""
    warning = ""
    if change_kind == "none":
        warning = (
            '<div class="warning">两版模型与提示词均无变化。若本次只是同一配置重跑，'
            "差异属于运行波动，不能作为版本优劣依据。</div>"
        )

    conclusion = data.get("conclusion")
    if conclusion == "unable":
        verdict = "两版差异在容差范围内或数据不可比，无法判定优劣"
    elif conclusion == "v2":
        verdict = f"{_esc(label2)} 效果更好"
    else:
        verdict = f"{_esc(label1)} 效果更好"

    version_cards: list[str] = []
    for index, (side_key, label) in enumerate((("v1", label1), ("v2", label2))):
        evaluation = evaluations[index] if index < len(evaluations) else {}
        success_rate = evaluation.get("success_rate")
        rate = _NA if not _is_number(success_rate) else f"{float(success_rate) * 100:.0f}%"
        models = ", ".join((object_bar.get("models") or {}).get(side_key) or []) or "--"
        prompt_nodes = ", ".join((object_bar.get("prompt_nodes") or {}).get(side_key) or []) or "--"
        version_cards.append(
            '<div class="version-card">'
            f'<div class="version-name">{_esc(label)}</div>'
            f'<div class="small">题目数 {evaluation.get("question_count") or 0} · 成功率 {rate} · '
            f"失败 {evaluation.get('failed_count') or 0}</div>"
            f"<div>模型：{_esc(models)}</div>"
            f"<div>提示词节点：{_esc(prompt_nodes)}</div>"
            "</div>"
        )

    cards = "".join(
        '<div class="card">'
        f'<div class="card-title">{_esc(metric.get("label"))}</div>'
        f'<div class="row"><span>{_esc(label1)}</span><b>{_fmt_metric_value(metric["key"], metric.get("v1"))}</b></div>'
        f'<div class="row"><span>{_esc(label2)}</span><b>{_fmt_metric_value(metric["key"], metric.get("v2"))}</b></div>'
        f'<div class="delta">{_fmt_metric_delta(metric["key"], metric.get("delta"))}</div>'
        "</div>"
        for metric in data.get("metrics") or []
    )

    per_question = data.get("per_question") or []
    grouped_bar = _svg_grouped_bar((data.get("charts") or {}).get("grouped_bar") or [], label1, label2)
    line_charts = "".join(
        _svg_line_chart(per_question, key, label, label1, label2) for key, label, _unit, _hib in METRIC_DEFS
    )
    detail = _html_detail_table(per_question, "v1", label1) + _html_detail_table(per_question, "v2", label2)
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")

    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8"/>
<title>评测对比报告 - {_esc(object_bar.get("app_name") or "--")}</title>
<style>{_EXPORT_HTML_CSS}</style>
</head>
<body>
<header>
  <h1>评测对比报告</h1>
  <div class="meta">生成时间：{generated_at}</div>
</header>
<section class="object-bar">
  <div class="object-line">应用：<b>{_esc(object_bar.get("app_name") or "--")}</b>
    测试集：<b>{_esc(object_bar.get("data_set") or "--")}</b>
    变更类型：<b>{_esc(change_label)}</b>{tag}</div>
  <div class="version-grid">{"".join(version_cards)}</div>
  {warning}
</section>
<h2>结论</h2>
<div class="conclusion-box">
  <div class="verdict">{verdict}</div>
  <div class="analysis"><span class="muted">对比分析：</span>{_esc(data.get("analysis"))}</div>
</div>
<h2>核心指标</h2>
<div class="cards">{cards}</div>
<h2>效果指标对比</h2>
{grouped_bar}
<h2>逐题趋势</h2>
<div class="line-grid">{line_charts}</div>
<h2>逐题明细</h2>
{detail}
</body>
</html>"""


def build_export_pdf(data: dict[str, Any]) -> bytes:
    """Render the full report to PDF via the project's Playwright/Chromium stack."""
    from playwright.sync_api import sync_playwright

    html_document = build_export_html(data)
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(args=["--no-sandbox", "--disable-dev-shm-usage"])
            try:
                page = browser.new_page()
                page.route("file://*", lambda route: route.abort())
                page.set_content(html_document, wait_until="load")
                return page.pdf(
                    format="A4",
                    margin={"top": "12mm", "right": "10mm", "bottom": "12mm", "left": "10mm"},
                    print_background=True,
                )
            finally:
                browser.close()
    except Exception:
        logger.exception("F090: PDF report rendering failed")
        raise


def _matplotlib_cjk_font():
    """Register a CJK font so chart titles/labels render in Chinese; None when absent."""
    from matplotlib import font_manager

    for family in (
        "WenQuanYi Zen Hei",
        "Noto Sans CJK SC",
        "Source Han Sans SC",
        "Microsoft YaHei",
        "SimHei",
        "PingFang SC",
    ):
        try:
            return font_manager.FontProperties(
                fname=font_manager.findfont(font_manager.FontProperties(family=family), fallback_to_default=False)
            )
        except Exception:
            continue
    return None


def _render_docx_charts(data: dict[str, Any], label1: str, label2: str) -> list[tuple[str, bytes]]:
    """Render the report charts to PNG bytes for embedding in the Word document."""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import PercentFormatter

    font = _matplotlib_cjk_font()
    if font is not None:
        from matplotlib import font_manager

        font_manager.fontManager.addfont(font.get_file())
        plt.rcParams["font.family"] = font.get_name()
    plt.rcParams["axes.unicode_minus"] = False

    charts: list[tuple[str, bytes]] = []

    grouped = (data.get("charts") or {}).get("grouped_bar") or []
    if grouped:
        figure, axis = plt.subplots(figsize=(8, 3.2), dpi=150)
        positions = list(range(len(grouped)))
        width = 0.35
        axis.bar(
            [x - width / 2 for x in positions],
            [float(m["v1"]) if _is_number(m.get("v1")) else 0 for m in grouped],
            width,
            label=label1,
            color="#94a3b8",
        )
        axis.bar(
            [x + width / 2 for x in positions],
            [float(m["v2"]) if _is_number(m.get("v2")) else 0 for m in grouped],
            width,
            label=label2,
            color="#2563eb",
        )
        axis.set_xticks(positions)
        axis.set_xticklabels([m.get("label") for m in grouped])
        axis.set_ylim(0, 1)
        axis.yaxis.set_major_formatter(PercentFormatter(1.0))
        axis.grid(axis="y", alpha=0.3)
        axis.legend()
        figure.tight_layout()
        buffer = io.BytesIO()
        figure.savefig(buffer, format="png")
        plt.close(figure)
        charts.append(("效果指标对比", buffer.getvalue()))

    per_question = data.get("per_question") or []
    if per_question:
        figure, axes = plt.subplots(3, 2, figsize=(10, 8), dpi=150)
        xs = list(range(1, len(per_question) + 1))
        for axis, (key, label, _unit, _hib) in zip(axes.flat, METRIC_DEFS):
            axis.plot(xs, _series_values(per_question, "v1", key), marker="o", label=label1, color="#94a3b8")
            axis.plot(xs, _series_values(per_question, "v2", key), marker="o", label=label2, color="#2563eb")
            axis.set_title(label)
            axis.set_xticks(xs)
            axis.grid(True, alpha=0.3)
            if key in _PERCENT_METRICS:
                axis.yaxis.set_major_formatter(PercentFormatter(1.0))
        axes.flat[0].legend()
        figure.tight_layout()
        buffer = io.BytesIO()
        figure.savefig(buffer, format="png")
        plt.close(figure)
        charts.append(("逐题趋势", buffer.getvalue()))
    return charts


def _docx_write_table(document, headers: list[str], rows: list[list[str]], font_size: float = 8) -> None:
    from docx.shared import Pt

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.size = Pt(font_size)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run("" if value is None else str(value))
            run.font.size = Pt(font_size)


def build_export_docx(data: dict[str, Any]) -> bytes:
    """Render the full comparison report as a Word (.docx) document."""
    from datetime import datetime

    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    evaluations = data.get("evaluations") or []
    label1 = _version_label(evaluations[0]) if len(evaluations) > 0 else "v1"
    label2 = _version_label(evaluations[1]) if len(evaluations) > 1 else "v2"
    object_bar = data.get("object_bar") or {}
    change_kind = object_bar.get("change_kind") or "none"
    change_label = _CHANGE_KIND_LABELS.get(change_kind, change_kind)

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    document.add_heading("评测对比报告", level=0)
    document.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    document.add_paragraph(f"应用：{object_bar.get('app_name') or '--'}")
    document.add_paragraph(f"测试集：{object_bar.get('data_set') or '--'}")
    document.add_paragraph(f"变更类型：{change_label}")
    if change_kind == "both":
        document.add_paragraph("多变量对比：模型与提示词同时变化，结果不可单独归因。")
    if change_kind == "none":
        document.add_paragraph(
            "两版模型与提示词均无变化。若本次只是同一配置重跑，差异属于运行波动，不能作为版本优劣依据。"
        )

    document.add_heading("对比版本", level=1)
    version_rows = [["版本", label1, label2]]
    for title, side_key in (("题目统计", None), ("模型", "models"), ("提示词节点", "prompt_nodes")):
        if side_key is None:
            values = []
            for index in range(2):
                evaluation = evaluations[index] if index < len(evaluations) else {}
                success_rate = evaluation.get("success_rate")
                rate = _NA if not _is_number(success_rate) else f"{float(success_rate) * 100:.0f}%"
                values.append(
                    f"题目数 {evaluation.get('question_count') or 0} · 成功率 {rate} · "
                    f"失败 {evaluation.get('failed_count') or 0}"
                )
        else:
            values = [
                ", ".join(object_bar.get(side_key, {}).get("v1") or []) or "--",
                ", ".join(object_bar.get(side_key, {}).get("v2") or []) or "--",
            ]
        version_rows.append([title, *values])
    _docx_write_table(document, ["项目", label1, label2], version_rows, font_size=9)

    conclusion = data.get("conclusion")
    if conclusion == "unable":
        verdict = "两版差异在容差范围内或数据不可比，无法判定优劣"
    elif conclusion == "v2":
        verdict = f"{label2} 效果更好"
    else:
        verdict = f"{label1} 效果更好"
    document.add_heading("结论", level=1)
    document.add_paragraph(verdict)
    document.add_paragraph("对比分析：")
    for line in str(data.get("analysis") or "").splitlines():
        if line.strip():
            document.add_paragraph(line)

    document.add_heading("核心指标", level=1)
    metric_rows = [
        [
            str(metric.get("label")),
            _fmt_metric_value(metric["key"], metric.get("v1")),
            _fmt_metric_value(metric["key"], metric.get("v2")),
            _fmt_metric_delta(metric["key"], metric.get("delta")),
        ]
        for metric in data.get("metrics") or []
    ]
    _docx_write_table(document, ["指标", label1, label2, "差值"], metric_rows, font_size=9)

    for title, png in _render_docx_charts(data, label1, label2):
        document.add_heading(title, level=1)
        document.add_picture(io.BytesIO(png), width=Inches(6.6))

    document.add_heading("逐题明细", level=1)
    per_question = data.get("per_question") or []
    for side_key, label in (("v1", label1), ("v2", label2)):
        document.add_heading(label, level=2)
        rows: list[list[str]] = []
        for index, entry in enumerate(per_question, start=1):
            side = entry.get(side_key) or {}
            status = _STATUS_LABELS.get(side.get("exec_status") or "", _NA)
            question = str(entry.get("question") or "")
            if entry.get("failed"):
                question += f"（失败：{entry.get('fail_reason') or '执行失败'}）"
            rows.append(
                [str(index), question, status]
                + [_export_cell(key, side.get(key)) for key, _label in _EXPORT_METRIC_LABELS]
            )
        _docx_write_table(
            document,
            ["#", "题目", "状态"] + [label for _key, label in _EXPORT_METRIC_LABELS],
            rows,
        )

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _safe_filename(base: str) -> str:
    import re

    cleaned = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", base).strip(" ._")
    return cleaned[:120] or "评测对比报告"


def _export_file_name(data: dict[str, Any], fmt: str) -> str:
    """Readable download name: app + report type + both version labels."""
    app_name = str((data.get("object_bar") or {}).get("app_name") or "").strip() or "评测对比"
    labels = [_version_label(e) for e in (data.get("evaluations") or [])[:2]]
    suffix = f"_{labels[0]}_vs_{labels[1]}" if len(labels) == 2 else ""
    return f"{_safe_filename(f'{app_name}评测报告{suffix}')}.{fmt}"


def build_compare_export(group_id: str, user_id: int, fmt: str = "xlsx") -> dict[str, str]:
    """Generate the two-version report (xlsx / html / pdf / docx), store it and return links."""
    from bisheng.core.storage.minio.minio_manager import get_minio_storage_sync
    from bisheng.utils import generate_uuid

    if fmt not in EXPORT_CONTENT_TYPES:
        raise ValueError(f"unsupported export format: {fmt}")

    data = build_compare_report(group_id, user_id)
    if fmt == "xlsx":
        content = build_export_xlsx(data)
    elif fmt == "html":
        content = build_export_html(data).encode("utf-8")
    elif fmt == "docx":
        content = build_export_docx(data)
    else:
        content = build_export_pdf(data)

    minio_client = get_minio_storage_sync()
    object_name = f"{_REPORT_DIR}/compare_{generate_uuid()}.{fmt}"
    minio_client.put_object_sync(
        bucket_name=minio_client.bucket,
        object_name=object_name,
        file=content,
        content_type=EXPORT_CONTENT_TYPES[fmt],
    )
    url = minio_client.get_share_link_sync(object_name, clear_host=False, expire_days=7)
    return {"url": url, "file_name": _export_file_name(data, fmt)}
