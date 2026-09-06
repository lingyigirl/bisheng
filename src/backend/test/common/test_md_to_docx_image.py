"""Regression tests for md_to_docx handling of <img> tags without src/title.

sanitize_html_for_pdf (also applied on the docx path) strips src attributes
whose URL is not allowed — e.g. leftover local /tmp paths from the parse
pipeline. DocxProcessor.add_picture must skip such <img> tags instead of
raising KeyError('src') and failing the whole export.
"""

import io

from docx import Document

from bisheng.common.utils.markdown_cmpnt.md_to_docx.markdocx import MarkDocx


def _render(md: str) -> bytes:
    docx_bytes, _ = MarkDocx()(md)
    return bytes(docx_bytes)


def _docx_text(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    return '\n'.join(p.text for p in doc.paragraphs)


def test_markdocx_skips_image_with_stripped_local_src() -> None:
    """An image pointing at a vanished local path loses its src in sanitizing
    and must be skipped, not raise KeyError('src')."""
    md = '# 现金流量表\n\n| 项目 | 金额 |\n| --- | --- |\n| 现金 | 100 |\n\n![image](/tmp/vanished/deadbeef.jpg)\n'
    docx_bytes = _render(md)
    assert docx_bytes
    assert '现金流量表' in _docx_text(docx_bytes)


def test_markdocx_skips_img_without_src_or_title() -> None:
    md = 'text before\n\n![]( )\n\ntext after'
    docx_bytes = _render(md)
    assert docx_bytes
    assert 'text before' in _docx_text(docx_bytes)
    assert 'text after' in _docx_text(docx_bytes)
