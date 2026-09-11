import io
import os
import tempfile
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from PIL import Image

from bisheng.workflow.nodes.report.docx_replace import DocxReplacer


class TestDocxReplacerFontSize(unittest.TestCase):
    """Regression: template run sizes (EMU Length, an int subclass) must not be
    re-applied as point values — that inflated 10.5pt runs to 133350pt."""

    def setUp(self):
        doc = Document()
        p = doc.add_paragraph()
        run = p.add_run("分析时间：")  # noqa: RUF001 (template uses Chinese punctuation)
        run.font.size = Pt(10.5)
        run = p.add_run("{{date}}")
        run.font.size = Pt(10.5)
        p = doc.add_paragraph()
        run = p.add_run("{{title}}")
        run.font.size = Pt(16)
        self.template = io.BytesIO()
        doc.save(self.template)
        self.template.seek(0)

    def test_template_font_sizes_survive_replacement(self):
        replacer = DocxReplacer(self.template)
        variables = {
            "date": [{"type": "text", "content": "2026-09-09"}],
            "title": [{"type": "text", "content": "章程分析"}],
        }
        with tempfile.TemporaryDirectory() as tmp_dir:
            output_path = os.path.join(tmp_dir, "out.docx")
            replacer.replace_and_save(variables, output_path)
            result = Document(output_path)

        sizes = []
        for paragraph in result.paragraphs:
            for run in paragraph.runs:
                if run.text.strip() and run.font.size is not None:
                    sizes.append(run.font.size.pt)

        self.assertTrue(sizes, "expected sized runs in the output document")
        self.assertLessEqual(max(sizes), 50, f"font size inflated: {sizes}")
        # static text segment keeps the template's 10.5pt
        texts = {run.text: run.font.size.pt if run.font.size else None
                 for paragraph in result.paragraphs for run in paragraph.runs}
        self.assertEqual(texts.get("分析时间："), 10.5)  # noqa: RUF001


class TestDocxReplacerImage(unittest.TestCase):
    """Image items: bytes content (e.g. decoded base64) renders; WEBP is converted to PNG."""

    def setUp(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("{{img}}")
        self.template = io.BytesIO()
        doc.save(self.template)
        self.template.seek(0)

    def _render(self, content):
        replacer = DocxReplacer(self.template)
        variables = {"img": [{"type": "image", "content": content}]}
        with tempfile.TemporaryDirectory() as tmp_dir:
            output_path = os.path.join(tmp_dir, "out.docx")
            replacer.replace_and_save(variables, output_path)
            return Document(output_path)

    def test_png_bytes_rendered(self):
        buf = io.BytesIO()
        Image.new("RGB", (4, 4), "red").save(buf, "PNG")
        result = self._render(buf.getvalue())
        self.assertGreaterEqual(len(result.inline_shapes), 1)
        self.assertNotIn("Image add failed", "\n".join(p.text for p in result.paragraphs))

    def test_webp_bytes_converted_to_png(self):
        buf = io.BytesIO()
        Image.new("RGB", (4, 4), "red").save(buf, "WEBP")
        result = self._render(buf.getvalue())
        self.assertGreaterEqual(len(result.inline_shapes), 1)
        self.assertNotIn("Image add failed", "\n".join(p.text for p in result.paragraphs))


class TestDocxReplacer(unittest.TestCase):
    def setUp(self):
        # need to prepare "test.png"、"test_chart1.png" in the same directory
        self.replacer = DocxReplacer("template.docx")

    def test_case(self):
        variables = {
            "var_1": [
                {"type": "text", "content": "这是正文 "},
                {"type": "text", "content": "这是粗体", "bold": True},
                {"type": "text", "content": " 继续正文"}
            ],
            "middle_placeholder": [
                {"type": "heading", "content": "数据表格", "level": 2},
                {
                    "type": "table",
                    "content": [
                        [
                            [
                                {"type": "text", "content": "列", "italic": True, "alignment": WD_ALIGN_PARAGRAPH.LEFT},
                                {"type": "text", "content": "1", "bold": True, "alignment": WD_ALIGN_PARAGRAPH.LEFT}
                            ],
                            {"type": "text", "content": "列2", "bold": True, "alignment": WD_ALIGN_PARAGRAPH.RIGHT}
                        ],
                        [
                            {"type": "text", "content": "数据1"},
                            {"type": "text", "content": "数据2"}
                        ]
                    ]
                }
            ],
            "report_date.123#1": [
                {"type": "text", "content": "2024-06-01"}
            ],
            "project_name": [
                {"type": "text", "content": "文档替换工具开发"}
            ],
            "executive_summary": [
                {"type": "text", "content": "本报告总结了文档替换工具的开发过程和使用方法。"},
                {"type": "image", "content": "test_chart1.png", "width": 5},
                {"type": "text", "content": "图表如上所示，展示了关键数据指标。"},
                {"type": "table", "content": [
                    [
                        {"type": "text", "content": "指标", "bold": True},
                        {"type": "text", "content": "数值", "bold": True},
                    ],
                    [
                        {"type": "text", "content": "指标A"},
                        {"type": "text", "content": "100"},
                    ], [
                        {"type": "text", "content": "指标B"},
                        {"type": "text", "content": "200"},
                    ],
                ]}
            ],
            "detailed_analysis": [
                {"type": "text", "content": "详细分析部分内容丰富，涵盖多个方面的数据解读。"},
                {"type": "heading", "content": "子标题：数据趋势", "level": 3},
                {"type": "text", "content": "通过对比历史数据，我们发现以下趋势..."}
            ],
            "chart_section": [
                {"type": "heading", "content": "红色图", "level": 3},
                {"type": "image", "content": "test.png"}
            ],
            "example_placeholder": [
                {"type": "text", "content": "这是一个示例占位符的替换内容。"}
            ],
            "table_title": [
                {"type": "text", "content": "以下是示例数据表格："}
            ],
            "table_date": [
                {"type": "text", "content": "2024-06-01"}
            ],
            "table_status": [
                {"type": "text", "content": "已完成"}
            ],
            "table_desc": [
                {"type": "text", "content": "这是表格的描述信息。"}
            ],
            "table_value": [
                {"type": "text", "content": "数值12345"}
            ],
            "table_summary": [
                {"type": "text", "content": "表格总结信息。"}
            ],
            "table_end": [
                {"type": "text", "content": "表格结束语。"}
            ],
            "only_placeholder": [
                {"type": "text", "content": "仅包含占位符的替换内容。"}
            ],
            "with_spaces": [
                {"type": "text", "content": "包含空格的    替换内容。"}
            ],
            "first": [
                {"type": "text", "content": "ONE"}
            ],
            "second": [
                {"type": "text", "content": "TWO"}
            ],
            "third": [
                {"type": "text", "content": "THREE"}
            ]
        }
        all_variables = self.replacer.extract_variables()
        print(all_variables)

        # 使用替换器
        self.replacer.replace_and_save(variables, "output.docx")
        print("文档替换完成！")


if __name__ == "__main__":
    unittest.main()
