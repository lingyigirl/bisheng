"""F090: report export — xlsx carries all six metrics per question; HTML/PDF are
full reports; the endpoint rejects unknown formats."""

import io

import pandas as pd
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from bisheng.evaluation.api.endpoints import evaluation as evaluation_endpoints
from bisheng.evaluation.domain.services.evaluation_compare_report import (
    _export_file_name,
    build_export_docx,
    build_export_html,
    build_export_pdf,
    build_export_xlsx,
    build_report_data,
)


def test_export_file_name_uses_app_and_versions():
    data = _report_data()
    data["object_bar"]["app_name"] = "流水分析智能体"
    assert _export_file_name(data, "pdf") == "流水分析智能体评测报告_v1_vs_v2.pdf"
    assert _export_file_name(data, "docx") == "流水分析智能体评测报告_v1_vs_v2.docx"


def test_export_file_name_sanitizes_illegal_characters():
    data = _report_data()
    data["object_bar"]["app_name"] = '流水/分析:*智能体?"'
    name = _export_file_name(data, "pdf")
    for char in '/\\:*?"<>|':
        assert char not in name


def test_docx_report_is_a_complete_document():
    from docx import Document as DocxDocument

    content = build_export_docx(_report_data())
    assert content[:2] == b"PK"  # docx is a zip container
    document = DocxDocument(io.BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for section in ("评测对比报告", "结论", "核心指标", "逐题明细"):
        assert section in text
    assert len(document.tables) >= 4  # versions + metrics + one detail table per side
    assert document.inline_shapes  # charts embedded


BASE = "/api/v1/evaluation"
_METRICS = ["准确率", "F1", "召回", "幻觉率", "成本", "延迟"]


def _row(q, *, precision, f1, recall, duration, tokens, contra=None, total=None, status="success", answer="a"):
    return {
        "question": q,
        "ground_truth": "g",
        "answer": answer,
        "recall": recall,
        "precision": precision,
        "F1": f1,
        "duration_ms": duration,
        "input_token": tokens // 2,
        "output_token": tokens - tokens // 2,
        "first_token_ms": None,
        "exec_status": status,
        "statements_contradicted_num": contra,
        "statements_num_answer": total,
    }


def _side(version, name, rows):
    return {
        "id": version,
        "version": version,
        "version_name": name,
        "app_name": "App",
        "data_set": "ds.csv",
        "status": 3,
        "result_score": None,
        "description": "",
        "flow_data": None,
        "rows": rows,
    }


_ROWS_V1 = [_row("q1", precision=0.5, f1=0.5, recall=0.5, duration=3000, tokens=30)]
_ROWS_V2 = [_row("q1", precision=0.9, f1=0.9, recall=0.9, duration=2000, tokens=20, contra=1, total=5)]


def _report_data():
    return build_report_data(_side(1, "v1", _ROWS_V1), _side(2, "v2", _ROWS_V2))


def test_export_has_two_sheets_readable_by_pandas():
    content = build_export_xlsx(_report_data())
    sheets = pd.read_excel(io.BytesIO(content), sheet_name=None)
    assert set(sheets) == {"汇总", "逐题明细"}


def test_detail_sheet_carries_all_six_metrics_for_both_versions():
    content = build_export_xlsx(_report_data())
    detail = pd.read_excel(io.BytesIO(content), sheet_name="逐题明细")
    for metric in _METRICS:
        assert f"{metric}(v1)" in detail.columns
        assert f"{metric}(v2)" in detail.columns
    row = detail.iloc[0]
    assert row["准确率(v1)"] == "50.00%"
    assert row["准确率(v2)"] == "90.00%"
    assert row["幻觉率(v2)"] == "20.00%"  # 1 / 5
    assert row["状态(v1)"] == "成功"


def test_export_is_full_and_independent_of_any_selected_metric():
    # the export input is the assembled report only — there is no metric selector
    data = _report_data()
    content = build_export_xlsx(data)
    assert isinstance(content, (bytes, bytearray)) and len(content) > 0


def test_html_report_is_a_complete_report():
    html = build_export_html(_report_data())
    for section in ("评测对比报告", "结论", "对比分析", "核心指标", "效果指标对比", "逐题趋势", "逐题明细"):
        assert section in html
    for metric in _METRICS:
        assert metric in html
    # full per-question values are present, not just the selected metric
    assert "50.00%" in html and "90.00%" in html
    assert "20.00%" in html  # hallucination 1 / 5


def test_html_report_escapes_dynamic_text():
    data = _report_data()
    data["object_bar"]["app_name"] = '<script>alert("x")</script>'
    data["per_question"][0]["question"] = "<img src=x onerror=alert(1)>"
    html = build_export_html(data)
    assert "<script>alert" not in html
    assert "<img src=x" not in html
    assert "&lt;script&gt;" in html


def test_pdf_report_renders_pdf_bytes():
    from playwright.sync_api import sync_playwright

    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(args=["--no-sandbox", "--disable-dev-shm-usage"])
            browser.close()
    except Exception as exc:
        pytest.skip(f"Playwright Chromium cannot launch in this environment: {exc}")

    content = build_export_pdf(_report_data())
    assert content.startswith(b"%PDF")
    assert len(content) > 1000


class _User:
    user_id = 7


async def _login():
    return _User()


@pytest.fixture
def client():
    from bisheng.common.dependencies.user_deps import UserPayload

    app = FastAPI()
    app.include_router(evaluation_endpoints.router, prefix="/api/v1")
    app.dependency_overrides[UserPayload.get_login_user] = _login
    return TestClient(app, raise_server_exceptions=False)


@pytest.mark.parametrize("fmt", ["xlsx", "html", "pdf", "docx"])
def test_export_endpoint_passes_format_and_returns_link(client, monkeypatch, fmt):
    calls = []

    def fake_export(gid, uid, selected):
        calls.append((gid, uid, selected))
        return {"url": f"http://minio/x.{selected}", "file_name": f"compare_{gid}.{selected}"}

    monkeypatch.setattr(evaluation_endpoints, "build_compare_export", fake_export)
    resp = client.get(f"{BASE}/compare/g1/report/export?format={fmt}")
    assert resp.status_code == 200
    assert resp.json()["data"]["url"] == f"http://minio/x.{fmt}"
    assert calls == [("g1", 7, fmt)]


def test_export_endpoint_rejects_unknown_format(client):
    resp = client.get(f"{BASE}/compare/g1/report/export?format=txt")
    assert resp.status_code == 400
