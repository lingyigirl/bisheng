"""覆盖 AC: AC-08-AC-19, AC-27: verify real report bytes, not mocked converters."""

import io
import zipfile

import pytest
from docx import Document
from PIL import Image

from bisheng.common.utils.report_document import render_docx, render_html
from bisheng.tool.domain.services.report_validation import validate_artifact


@pytest.fixture
def image_bytes():
    stream = io.BytesIO()
    Image.new("RGB", (320, 120), "navy").save(stream, format="PNG")
    return stream.getvalue()


def test_real_docx_embeds_chinese_path_image_and_restarts_lists(image_bytes):
    source = "# Report\n\n1. First\n2. Second\n\n## Another\n\n1. Restart\n\n![Chart](charts/中文 图.png)\n\n|A|B|\n|---|---|\n|one|long text|"
    result = render_docx(source, {"charts/中文 图.png": image_bytes})
    doc = Document(io.BytesIO(result))
    assert len(doc.inline_shapes) == 1
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    num_ids = [p._p.xpath("./w:pPr/w:numPr/w:numId/@w:val") for p in doc.paragraphs]
    ids = [x[0] for x in num_ids if x]
    assert ids[0] == ids[1] and ids[2] != ids[0]
    assert validate_artifact("report.docx", result)["images"] == 1


def test_missing_or_invalid_image_is_an_explicit_failure(image_bytes):
    with pytest.raises(ValueError, match=r"missing\.png"):
        render_docx("![Missing](missing.png)", {})
    with pytest.raises(ValueError):
        render_docx("![Bad](bad.png)", {"bad.png": b"not an image"})


def test_loose_list_keeps_number_and_bold_lead_on_same_paragraph():
    source = "1. **First conclusion.** Body text.\n\n2. **Second conclusion.** More text."
    doc = Document(io.BytesIO(render_docx(source, {})))
    assert [p.text for p in doc.paragraphs] == ["First conclusion. Body text.", "Second conclusion. More text."]
    assert all(p._p.xpath("./w:pPr/w:numPr") for p in doc.paragraphs)
    assert all(p.runs[0].bold for p in doc.paragraphs)


def test_list_continuations_and_nested_items_preserve_order_without_extra_numbers():
    source = "1. **Lead.** Body.\n\n    Continuation.\n\n    - Nested.\n\n    After nested.\n\n2. Next."
    doc = Document(io.BytesIO(render_docx(source, {})))
    assert [p.text for p in doc.paragraphs] == ["Lead. Body.", "Continuation.", "Nested.", "After nested.", "Next."]
    assert [bool(p._p.xpath("./w:pPr/w:numPr")) for p in doc.paragraphs] == [True, False, True, False, True]


def test_pdf_manifest_counts_shared_raster_images_and_bounds_sample_pages(image_bytes):
    import fitz

    with fitz.open() as doc:
        image_id = 0
        for _ in range(5):
            page = doc.new_page()
            image_id = page.insert_image(fitz.Rect(72, 72, 232, 132), stream=image_bytes, xref=image_id)
        result = validate_artifact("report.pdf", doc.tobytes())
    assert result["raster_images"] == 1
    assert result["image_pages"] == [1, 2, 3]
    assert result["pages"] == 5


def test_pdf_image_pages_do_not_include_unused_shared_resources(image_bytes):
    import fitz

    with fitz.open() as doc:
        first = doc.new_page()
        first.insert_image(fitz.Rect(72, 72, 232, 132), stream=image_bytes)
        first_xref = first.xref
        second = doc.new_page()
        resources = doc.xref_get_key(first_xref, "Resources")[1]
        doc.xref_set_key(second.xref, "Resources", resources)
        assert second.get_images() and not second.get_image_info()
        result = validate_artifact("report.pdf", doc.tobytes())
    assert result["raster_images"] == 1
    assert result["image_pages"] == [1]


def test_html_contains_embedded_resources_and_no_external_scripts(image_bytes):
    html = render_html("# Report\n\n![Chart](chart.png)\n<script>alert(1)</script>", {"chart.png": image_bytes})
    assert "data:image/png;base64," in html
    assert "<script" not in html
    assert "table-header-group" in html


def test_fake_office_and_pdf_files_cannot_pass_validation():
    for name in ["a.docx", "a.pdf", "a.xlsx", "a.pptx", "a.png"]:
        with pytest.raises(ValueError):
            validate_artifact(name, b'{"error":"conversion failed"}')


def test_zip_expansion_is_bounded():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("huge.xml", b"0" * 2000)
    with pytest.raises(ValueError, match="expanded"):
        validate_artifact("a.docx", stream.getvalue(), max_expanded_bytes=1000)


def test_real_spreadsheet_presentation_csv_and_pdf():
    import fitz
    from openpyxl import Workbook
    from pptx import Presentation

    stream = io.BytesIO()
    workbook = Workbook()
    workbook.active.append(["Item", "Value"])
    workbook.active.append(["A", 42])
    workbook.save(stream)
    assert validate_artifact("a.xlsx", stream.getvalue())["sheets"] == 1
    stream = io.BytesIO()
    slides = Presentation()
    slides.slides.add_slide(slides.slide_layouts[0])
    slides.save(stream)
    assert validate_artifact("a.pptx", stream.getvalue())["slides"] == 1
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Verified PDF")
    assert validate_artifact("a.pdf", doc.tobytes())["pages"] == 1
    doc.close()
    assert validate_artifact("a.csv", b"item,value\na,42\n")["rows"] == 2


def test_spreadsheet_corrupt_sheet_cannot_pass_lazy_loading():
    from openpyxl import Workbook

    original, broken = io.BytesIO(), io.BytesIO()
    Workbook().save(original)
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(broken, "w") as target:
        for name in source.namelist():
            target.writestr(name, b"<broken" if name == "xl/worksheets/sheet1.xml" else source.read(name))
    with pytest.raises(ValueError, match=r"Invalid \.xlsx"):
        validate_artifact("broken.xlsx", broken.getvalue())


async def test_artifact_revalidates_missing_resource_after_success(tmp_path, image_bytes, monkeypatch):
    from bisheng.linsight.domain.services import artifact_validation
    from bisheng.linsight.domain.services.artifact_validation import validate_deliverable
    from bisheng.tool.domain.services.report_runtime import _no_capacity

    monkeypatch.setattr(artifact_validation, "report_capacity", _no_capacity)

    report = tmp_path / "report.md"
    report.write_text("![Chart](chart.png)", encoding="utf-8")
    (tmp_path / "chart.png").write_bytes(image_bytes)
    assert (await validate_deliverable(str(report), str(tmp_path)))["resources"] == ["chart.png"]
    (tmp_path / "chart.png").unlink()
    with pytest.raises(ValueError, match="Missing"):
        await validate_deliverable(str(report), str(tmp_path))
